"""
anomaly_pipeline.py
===================
Fully automatic anomaly-detection pipeline, refactored from the
"Vehicle Anomaly Detection System" notebook.

    upload (bytes) -> load/validate -> profile -> clean numeric text
    -> detect dates -> plan feature columns -> EDA -> preprocess
    -> train autoencoder -> score rows -> explain -> export

Everything is exposed through `run_pipeline(...)`, which returns a
`PipelineResult` that the Streamlit UI (app.py) renders. The module has no
Streamlit dependency, so it can also be used from a script or a notebook.
"""

from __future__ import annotations

import io
import json
import re
import tempfile
import time
import warnings
import zipfile
from dataclasses import dataclass, field
from itertools import combinations
from pathlib import Path
from typing import Callable, Optional

import joblib
import matplotlib

matplotlib.use("Agg")
from matplotlib.figure import Figure  # noqa: E402
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
import seaborn as sns  # noqa: E402
from sklearn.compose import ColumnTransformer  # noqa: E402
from sklearn.impute import SimpleImputer  # noqa: E402
from sklearn.model_selection import train_test_split  # noqa: E402
from sklearn.pipeline import Pipeline  # noqa: E402
from sklearn.preprocessing import OneHotEncoder, StandardScaler  # noqa: E402

__version__ = "1.0.0"

ProgressFn = Callable[[float, str], None]

PRIMARY = "#6366f1"
DANGER = "#ef4444"
MUTED = "#94a3b8"

ILLEGAL_XLSX_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


# ----------------------------------------------------------------------
# Configuration and result containers
# ----------------------------------------------------------------------
@dataclass
class PipelineConfig:
    test_size: float = 0.20
    random_state: int = 42
    threshold_percentile: float = 95.0
    epochs: int = 50
    batch_size: int = 32
    patience: int = 5
    validation_split: float = 0.10
    numeric_text_threshold: float = 0.95
    mostly_numeric_threshold: float = 0.90
    identifier_unique_ratio: float = 0.90
    max_onehot_categories: int = 30
    min_category_frequency: float = 0.01
    score_all_rows: bool = True
    dayfirst: bool = False
    max_eda_columns: int = 12
    min_rows: int = 30


@dataclass
class PipelineResult:
    filename: str
    config: PipelineConfig
    original: pd.DataFrame
    scored: pd.DataFrame
    anomalies: pd.DataFrame
    summary: dict
    tables: dict = field(default_factory=dict)
    figures: dict = field(default_factory=dict)
    history: dict = field(default_factory=dict)
    warnings: list = field(default_factory=list)
    excel_report: bytes = b""
    artifacts_zip: bytes = b""


# ----------------------------------------------------------------------
# 1. Loading and validation
# ----------------------------------------------------------------------
def list_excel_sheets(file_bytes: bytes, filename: str) -> list:
    """Return sheet names for an Excel upload (empty list for other types)."""
    if Path(filename).suffix.lower() not in {".xlsx", ".xlsm", ".xls"}:
        return []
    try:
        return list(pd.ExcelFile(io.BytesIO(file_bytes)).sheet_names)
    except Exception:
        return []


def _detect_delimiter(text_sample: str) -> str:
    """Pick among , ; tab | only (pandas' own sniffer can choose '.' or letters)."""
    import csv

    head = "\n".join(text_sample.splitlines()[:20])
    try:
        return csv.Sniffer().sniff(head, delimiters=",;\t|").delimiter
    except csv.Error:
        return ","


def load_dataset(file_bytes: bytes, filename: str, sheet_name=None) -> pd.DataFrame:
    """Load and validate a CSV / Excel upload. Returns a clean-named DataFrame."""
    if not file_bytes:
        raise ValueError("The uploaded file is empty.")

    ext = Path(filename).suffix.lower()
    buf = io.BytesIO(file_bytes)

    if ext in {".csv", ".tsv", ".txt"}:
        df, last_error = None, None
        for encoding in ("utf-8-sig", "latin-1"):
            sample = file_bytes[:65536].decode(encoding, errors="ignore")
            sep = _detect_delimiter(sample)
            try:
                buf.seek(0)
                df = pd.read_csv(buf, sep=sep, encoding=encoding)
                break
            except Exception as exc:  # noqa: BLE001
                last_error = exc
        if df is None:
            raise ValueError(f"Could not read the CSV file: {last_error}")

    elif ext in {".xlsx", ".xlsm"}:
        df = pd.read_excel(buf, sheet_name=sheet_name or 0, engine="openpyxl")

    elif ext == ".xls":
        try:
            df = pd.read_excel(buf, sheet_name=sheet_name or 0, engine="xlrd")
        except ImportError as exc:
            raise ValueError(
                "Reading .xls files needs the 'xlrd' package "
                "(pip install xlrd), or save the file as .xlsx."
            ) from exc
    else:
        raise ValueError("Unsupported file format. Upload a CSV, XLSX or XLS file.")

    # Excel exports often carry trailing blank rows/columns.
    df = df.dropna(how="all").dropna(axis=1, how="all").reset_index(drop=True)

    if df.empty:
        raise ValueError("The uploaded dataset contains no rows.")
    if df.shape[1] == 0:
        raise ValueError("The uploaded dataset contains no columns.")

    # Tidy column names: strip, fill blanks, make unique.
    names, seen = [], {}
    for position, col in enumerate(df.columns):
        name = str(col).strip() or f"column_{position + 1}"
        if name in seen:
            seen[name] += 1
            name = f"{name}.{seen[name]}"
        else:
            seen[name] = 0
        names.append(name)
    df.columns = names

    return df


# ----------------------------------------------------------------------
# 2. Profiling
# ----------------------------------------------------------------------
def profile_dataset(df: pd.DataFrame) -> pd.DataFrame:
    profile = pd.DataFrame(
        {
            "column": df.columns.astype(str),
            "data_type": [str(t) for t in df.dtypes],
            "missing_count": df.isna().sum().values,
            "missing_percentage": (df.isna().mean().values * 100).round(2),
            "unique_values": [df[c].nunique(dropna=True) for c in df.columns],
        }
    )
    return profile.sort_values("missing_percentage", ascending=False).reset_index(drop=True)


def _normalize_for_comparison(series: pd.Series) -> pd.Series:
    result = series.astype("string").str.strip().str.casefold()
    return result.str.replace(r"\s+", " ", regex=True)


def compare_columns(df: pd.DataFrame, max_columns: int = 60, sample_rows: int = 5000) -> pd.DataFrame:
    """Report column pairs whose (normalised) values (almost) always agree."""
    columns = list(df.columns)[:max_columns]
    sample = df[columns]
    if len(sample) > sample_rows:
        sample = sample.sample(sample_rows, random_state=0)

    normalized = {c: _normalize_for_comparison(sample[c]) for c in columns}
    rows = []
    for a, b in combinations(columns, 2):
        both = normalized[a].notna() & normalized[b].notna()
        n = int(both.sum())
        if n == 0:
            continue
        rate = float(normalized[a][both].eq(normalized[b][both]).mean() * 100)
        if rate >= 95:
            rows.append(
                {
                    "column_a": a,
                    "column_b": b,
                    "match_rate_pct": round(rate, 2),
                    "comparable_rows": n,
                }
            )
    return pd.DataFrame(
        rows, columns=["column_a", "column_b", "match_rate_pct", "comparable_rows"]
    ).sort_values("match_rate_pct", ascending=False, ignore_index=True)


# ----------------------------------------------------------------------
# 3. Temporal column detection
# ----------------------------------------------------------------------
_MONTH = (
    r"(?:jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|jun(?:e)?|"
    r"jul(?:y)?|aug(?:ust)?|sep(?:tember)?|oct(?:ober)?|nov(?:ember)?|dec(?:ember)?)"
)
_MONTH_YEAR_RE = re.compile(
    rf"^(?:{_MONTH}\s+\d{{4}}|\d{{4}}[-/]\d{{1,2}}|\d{{1,2}}[-/]\d{{4}})$", re.IGNORECASE
)
_FULL_DATE_RE = re.compile(r"^(?:\d{4}[-/]\d{1,2}[-/]\d{1,2}|\d{1,2}[-/]\d{1,2}[-/]\d{4})$")


def detect_temporal_columns(
    df: pd.DataFrame,
    year_min: int = 1900,
    year_max: Optional[int] = None,
    threshold: float = 0.80,
    sample_size: int = 5000,
    dayfirst: bool = False,
) -> pd.DataFrame:
    """Detect full-date, month-year and year-only columns."""
    year_max = year_max or pd.Timestamp.now().year + 1
    results = []

    for column in df.columns:
        series = df[column]

        if pd.api.types.is_datetime64_any_dtype(series):
            results.append(
                {
                    "column": column,
                    "temporal_type": "full_date",
                    "matched_percentage": 100.0,
                    "missing_count": int(series.isna().sum()),
                }
            )
            continue

        values = series.dropna()
        if values.empty or pd.api.types.is_bool_dtype(series):
            continue
        if len(values) > sample_size:
            values = values.sample(sample_size, random_state=0)

        if pd.api.types.is_numeric_dtype(series):
            numeric = pd.to_numeric(values, errors="coerce")
            valid = numeric.between(year_min, year_max) & (numeric % 1 == 0)
            pct = 100 * valid.mean()
            if pct >= threshold * 100 and numeric.nunique() > 1:
                results.append(
                    {
                        "column": column,
                        "temporal_type": "year_only",
                        "matched_percentage": round(pct, 2),
                        "missing_count": int(series.isna().sum()),
                    }
                )
            continue

        text = values.astype(str).str.strip()
        year_only = text.str.fullmatch(r"\d{4}", na=False)
        years = pd.to_numeric(text.where(year_only), errors="coerce")
        valid_year = year_only & years.between(year_min, year_max)
        month_year = text.map(lambda v: bool(_MONTH_YEAR_RE.fullmatch(v)))
        full_date = text.map(lambda v: bool(_FULL_DATE_RE.fullmatch(v)))

        valid_full = pd.Series(False, index=text.index)
        if full_date.any():
            valid_full.loc[full_date] = pd.to_datetime(
                text.loc[full_date], errors="coerce", format="mixed", dayfirst=dayfirst
            ).notna()
        valid_my = pd.Series(False, index=text.index)
        if month_year.any():
            valid_my.loc[month_year] = pd.to_datetime(
                text.loc[month_year], errors="coerce", format="mixed", dayfirst=dayfirst
            ).notna()

        categories = pd.DataFrame(
            {"year_only": valid_year, "month_year": valid_my, "full_date": valid_full}
        )
        pct = 100 * categories.any(axis=1).mean()
        if pct < threshold * 100:
            continue

        found = categories.columns[categories.any(axis=0)].tolist()
        results.append(
            {
                "column": column,
                "temporal_type": found[0] if len(found) == 1 else "mixed_" + "_and_".join(found),
                "matched_percentage": round(pct, 2),
                "missing_count": int(series.isna().sum()),
            }
        )

    return pd.DataFrame(
        results, columns=["column", "temporal_type", "matched_percentage", "missing_count"]
    )


# ----------------------------------------------------------------------
# 4. Numeric-text conversion (prices, engine size, power, torque, ...)
# ----------------------------------------------------------------------
def parse_alternate_measurement(value, measurement_type: str):
    """
    Parse explicitly labelled automotive power or torque values.
    power -> kW, torque -> Nm. Returns (numeric_value, status).
    """
    if measurement_type not in {"power", "torque"}:
        raise ValueError("measurement_type must be 'power' or 'torque'")

    if value is None:
        return None, "missing"
    try:
        if pd.isna(value):
            return None, "missing"
    except (TypeError, ValueError):
        pass

    text = str(value).strip().lower().replace(",", "")
    if not text:
        return None, "missing"

    if measurement_type == "power":
        standard_unit = "kw"
        factors = {"kw": 1.0, "bhp": 0.745699872, "hp": 0.745699872, "ps": 0.73549875}
        unit_pattern = r"bhp|kw|ps|hp"
    else:
        standard_unit = "nm"
        factors = {"nm": 1.0, "kgm": 9.80665, "kgf-m": 9.80665}
        unit_pattern = r"kgf[\s-]*m|kgm|nm"

    # Normalise repeated suffixes, e.g. 402.30bhpbhp.
    text = re.sub(rf"({unit_pattern})\s*\1\b", r"\1", text)

    measurement = re.compile(
        rf"(?<![\w.])(-?\d+(?:\.\d+)?)\s*({unit_pattern})"
        rf"(?=$|[\s()\[\]/,;@]|at\b|\d)",
        re.IGNORECASE,
    )
    matches = [(float(m.group(1)), m.group(2).lower()) for m in measurement.finditer(text)]
    if not matches:
        return None, "unresolved"

    standard = [n for n, u in matches if u == standard_unit]
    if len(standard) == 1:
        return standard[0], "parsed"
    if len(standard) > 1:
        return None, "unresolved"

    number, unit = matches[0]
    unit = re.sub(r"[\s-]+", "", unit)
    if unit == "kgfm":
        unit = "kgf-m"
    return number * factors[unit], "parsed"


_PRICE_RE = re.compile(
    r"^(?:₹|rs\.?|inr)?\s*(\d+(?:,\d{2,3})*(?:\.\d+)?)\s*"
    r"(lakh|lakhs|lac|lacs|crore|crores|cr)?$",
    re.IGNORECASE,
)
_NUMBER_RE = re.compile(r"^[+-]?(?:\d{1,3}(?:,\d{2,3})+|\d+)(?:\.\d+)?$")
_GENERIC_UNIT_RE = re.compile(
    r"^([+-]?\d+(?:\.\d+)?)\s*(kmpl|km/kg|kms?|cc|rpm|seats?|inch|inches|mm|kg)$",
    re.IGNORECASE,
)
_LABELLED_UNIT_RE = re.compile(r"(-?\d+(?:\.\d+)?)\s*(kgf[\s-]*m|kgm|bhp|ps|hp|kw|nm)", re.IGNORECASE)


def _parse_value(value, is_price: bool = False):
    """Return (numeric_value, status, standard_unit)."""
    if value is None:
        return np.nan, "missing", None
    try:
        if pd.isna(value):
            return np.nan, "missing", None
    except (TypeError, ValueError):
        pass

    text = str(value).strip()
    if not text:
        return None, "unrecognized", None

    normalized = re.sub(r"\s+", " ", text.replace(",", "")).strip()

    if is_price:
        match = _PRICE_RE.fullmatch(text)
        if not match:
            return None, "unrecognized", None
        number = float(match.group(1).replace(",", ""))
        suffix = (match.group(2) or "").lower()
        if suffix in {"lakh", "lakhs", "lac", "lacs"}:
            number *= 100_000
        elif suffix in {"crore", "crores", "cr"}:
            number *= 10_000_000
        return number, "parsed", None

    if _NUMBER_RE.fullmatch(normalized):
        return float(normalized.replace(",", "")), "parsed", None

    generic = _GENERIC_UNIT_RE.fullmatch(normalized)
    if generic:
        return float(generic.group(1)), "parsed", generic.group(2).lower()

    cleaned = re.sub(
        r"\b(bhp|kw|ps|hp|nm|kgm|kgf[\s-]*m)\s*\1\b", r"\1", normalized, flags=re.IGNORECASE
    )
    found = list(_LABELLED_UNIT_RE.finditer(cleaned))
    if found:
        units = {m.group(2).lower().replace(" ", "") for m in found}
        has_power = bool(units & {"bhp", "hp", "ps", "kw"})
        has_torque = bool(units & {"nm", "kgm", "kgf-m"})
        if has_power and has_torque:
            return None, "unrecognized", None
        if has_power:
            number, status = parse_alternate_measurement(cleaned, "power")
            return (number, "parsed", "kW") if status == "parsed" else (None, "unrecognized", None)
        if has_torque:
            number, status = parse_alternate_measurement(cleaned, "torque")
            return (number, "parsed", "Nm") if status == "parsed" else (None, "unrecognized", None)

    return None, "unrecognized", None


def convert_numeric_text_columns(df: pd.DataFrame, temporal_report: pd.DataFrame, threshold: float = 0.95):
    """
    Convert consistently numeric text columns into numbers.
    Does not modify the input. Returns (converted_df, conversion_log, review_report).
    """
    out = df.copy(deep=True)
    protected = set(temporal_report["column"]) if len(temporal_report) else set()
    conversion_records, review_records = [], []

    for column in out.columns:
        series = out[column]
        if column in protected:
            continue
        if not (pd.api.types.is_object_dtype(series) or pd.api.types.is_string_dtype(series)):
            continue
        non_null = series.dropna()
        if non_null.empty:
            continue

        is_price = "price" in str(column).lower()

        # Fast pre-check on a sample so free-text columns are skipped cheaply.
        probe = non_null.sample(min(500, len(non_null)), random_state=0)
        probe_ratio = np.mean([_parse_value(v, is_price)[1] == "parsed" for v in probe])
        if probe_ratio < max(0.0, threshold - 0.15):
            continue

        parsed = [_parse_value(v, is_price) if pd.notna(v) else (np.nan, "missing", None) for v in series.tolist()]
        statuses = [p[1] for p in parsed]
        recognized = sum(s == "parsed" for s in statuses)
        ratio = recognized / len(non_null)
        if ratio < threshold:
            continue

        values = series.tolist()
        output, unresolved_positions = [], []
        for position, (number, status, _unit) in enumerate(parsed):
            if status == "parsed":
                output.append(number)
            elif status == "missing":
                output.append(np.nan)
            else:
                output.append(values[position])
                unresolved_positions.append(position)

        for position in unresolved_positions:
            review_records.append(
                {
                    "row_number": position + 1,
                    "column": column,
                    "original_value": values[position],
                    "reason": "Unrecognized or ambiguous format",
                }
            )

        out[column] = pd.Series(output, index=series.index, dtype=object)
        if not unresolved_positions:
            out[column] = pd.to_numeric(out[column], errors="raise")

        units = {p[2] for p in parsed if p[1] == "parsed" and p[2] is not None}
        conversion_records.append(
            {
                "column": column,
                "recognized_numeric_percentage": round(ratio * 100, 2),
                "original_dtype": str(series.dtype),
                "converted_dtype": str(out[column].dtype),
                "unresolved_values": len(unresolved_positions),
                "standard_unit": next(iter(units)) if len(units) == 1 else None,
            }
        )

    conversion_log = pd.DataFrame(
        conversion_records,
        columns=[
            "column",
            "recognized_numeric_percentage",
            "original_dtype",
            "converted_dtype",
            "unresolved_values",
            "standard_unit",
        ],
    )
    review_report = pd.DataFrame(
        review_records, columns=["row_number", "column", "original_value", "reason"]
    )
    return out, conversion_log, review_report


# ----------------------------------------------------------------------
# 5. Feature-column planning
# ----------------------------------------------------------------------
def _is_identifier_name(name: str) -> bool:
    text = str(name).strip()
    spaced = re.sub(r"([a-z])([A-Z])", r"\1_\2", text)
    normalized = re.sub(r"[^A-Za-z0-9]+", "_", spaced).lower().strip("_")
    tokens = normalized.split("_") if normalized else [""]
    return (
        normalized.startswith("unnamed")
        or tokens in (["id"], ["index"], ["uuid"])
        or tokens[-1:] in (["id"], ["identifier"], ["uuid"])
    )


def plan_columns(df: pd.DataFrame, temporal_report: pd.DataFrame, cfg: PipelineConfig) -> pd.DataFrame:
    """
    Decide, for every column, whether it is a model feature
    (numeric / categorical / date) or excluded, and why.
    Nothing is deleted from the source data.
    """
    temporal_types = (
        dict(zip(temporal_report["column"], temporal_report["temporal_type"]))
        if len(temporal_report)
        else {}
    )
    rows = []
    seen_hashes = {}

    for col in df.columns:
        series = df[col]
        non_null = series.dropna()
        n_unique = int(non_null.nunique()) if len(non_null) else 0
        record = {
            "column": col,
            "dtype": str(series.dtype),
            "role": "excluded",
            "reason": "",
            "missing_percentage": round(float(series.isna().mean() * 100), 2),
            "unique_values": n_unique,
        }

        def done(role, reason):
            record["role"], record["reason"] = role, reason
            rows.append(record)

        if non_null.empty:
            done("excluded", "Completely empty")
            continue
        if n_unique <= 1:
            done("excluded", "Constant column (single value)")
            continue

        # Exact duplicate of an earlier column (first one is kept).
        try:
            digest = pd.util.hash_pandas_object(series, index=False).values.tobytes()
            first = seen_hashes.get(digest)
            if first is not None and df[first].equals(series):
                done("excluded", f"Exact duplicate of '{first}'")
                continue
            seen_hashes.setdefault(digest, col)
        except TypeError:
            pass  # unhashable (list-like) values are handled below

        if col in temporal_types and not (
            pd.api.types.is_numeric_dtype(series) and temporal_types[col] == "year_only"
        ):
            done("date", f"Date column ({temporal_types[col]}) expanded into numeric parts")
            continue

        sample = non_null.head(1000)
        if pd.api.types.is_object_dtype(series) and sample.map(
            lambda v: isinstance(v, (list, tuple, set, dict))
        ).any():
            done("excluded", "Contains list-like values")
            continue

        if _is_identifier_name(col):
            done("excluded", "Identifier / index column (by name)")
            continue

        is_bool = pd.api.types.is_bool_dtype(series)
        if is_bool:
            done("categorical", "Boolean feature (one-hot encoded)")
            continue
        is_numeric = pd.api.types.is_numeric_dtype(series)

        if is_numeric:
            numeric = pd.to_numeric(non_null, errors="coerce")
            if (
                len(numeric) >= 10
                and numeric.nunique() == len(numeric)
                and (numeric % 1 == 0).all()
                and len(series) == len(non_null)
            ):
                diffs = numeric.reset_index(drop=True).diff().dropna()
                if diffs.eq(1).all() or diffs.eq(-1).all():
                    done("excluded", "Sequential index column")
                    continue
            done("numeric", "Numeric feature")
            continue

        # Text-like column.
        ratio = n_unique / len(non_null)
        numeric_share = float(pd.to_numeric(non_null, errors="coerce").notna().mean())
        if numeric_share >= cfg.mostly_numeric_threshold:
            done("numeric", "Mostly numeric text (unparseable values become missing)")
            continue
        if ratio >= cfg.identifier_unique_ratio:
            done("excluded", "Identifier-like text (almost every value is unique)")
            continue
        if n_unique > 100 and ratio > 0.5:
            done("excluded", "High-cardinality text")
            continue
        done("categorical", "Categorical feature (one-hot encoded)")

    return pd.DataFrame(rows)


def _clean_categorical(series: pd.Series) -> pd.Series:
    """Strings with a real NaN for missing, so encoders never meet mixed types."""
    out = series.astype("string").astype(object)
    return out.where(series.notna(), np.nan)


def _parse_dates(series: pd.Series, dayfirst: bool) -> pd.Series:
    if pd.api.types.is_datetime64_any_dtype(series):
        return series
    return pd.to_datetime(series.astype("string"), errors="coerce", format="mixed", dayfirst=dayfirst)


def build_model_frame(df: pd.DataFrame, plan: pd.DataFrame, temporal_report: pd.DataFrame, cfg: PipelineConfig):
    """
    Build the model-ready feature frame from the planned columns.
    Returns (model_df, numeric_features, categorical_features, date_info).
    """
    temporal_types = (
        dict(zip(temporal_report["column"], temporal_report["temporal_type"]))
        if len(temporal_report)
        else {}
    )
    frame, numeric, categorical, date_info = {}, [], [], {}

    for _, row in plan.iterrows():
        col, role = row["column"], row["role"]
        if role == "numeric":
            values = pd.to_numeric(df[col], errors="coerce").astype("float64")
            frame[col] = values.replace([np.inf, -np.inf], np.nan)
            numeric.append(col)
        elif role == "categorical":
            frame[col] = _clean_categorical(df[col])
            categorical.append(col)
        elif role == "date":
            kind = temporal_types.get(col, "full_date")
            series = df[col]
            if kind == "year_only":
                parts = {col: pd.to_numeric(series, errors="coerce")}
            else:
                dates = _parse_dates(series, cfg.dayfirst)
                parts = {f"{col}__year": dates.dt.year, f"{col}__month": dates.dt.month}
                if "full_date" in kind:
                    parts[f"{col}__dayofweek"] = dates.dt.dayofweek
            created = []
            for name, values in parts.items():
                values = pd.to_numeric(values, errors="coerce")
                if values.dropna().nunique() > 1:
                    frame[name] = values.astype("float64")
                    numeric.append(name)
                    created.append(name)
            date_info[col] = created

    model_df = pd.DataFrame(frame, index=df.index)
    return model_df, numeric, categorical, date_info


# ----------------------------------------------------------------------
# 6. Excel-safe export helpers
# ----------------------------------------------------------------------
def _excel_safe(df: pd.DataFrame) -> pd.DataFrame:
    safe = df.copy()
    for col in safe.columns:
        if isinstance(safe[col].dtype, pd.DatetimeTZDtype):
            safe[col] = safe[col].dt.tz_localize(None)
        elif safe[col].dtype == object or pd.api.types.is_string_dtype(safe[col]):
            safe[col] = safe[col].map(
                lambda v: ILLEGAL_XLSX_CHARS.sub("", v)
                if isinstance(v, str)
                else (str(v) if isinstance(v, (list, tuple, set, dict)) else v)
            )
    return safe


def _style_sheet(ws, df: pd.DataFrame):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    ws.freeze_panes = "A2"
    if ws.max_row > 1 and ws.max_column > 0:
        ws.auto_filter.ref = ws.dimensions
    fill = PatternFill("solid", fgColor="4F46E5")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(vertical="center")
    for i, col in enumerate(df.columns, start=1):
        sample = df[col].head(200).astype(str).str.len().max() if len(df) else 0
        width = min(max(len(str(col)), int(sample or 0)) + 2, 45)
        ws.column_dimensions[get_column_letter(i)].width = max(width, 10)


def build_excel_report(result: PipelineResult) -> bytes:
    sheets = {
        "Summary": pd.DataFrame(
            [{"Metric": k, "Value": str(v)} for k, v in result.summary.items()]
        ),
        "Anomalies": result.anomalies,
    }
    if len(result.scored) <= 1_000_000:
        sheets["All_Rows_Scored"] = result.scored
    sheets["Column_Plan"] = result.tables.get("column_plan", pd.DataFrame())
    sheets["Conversion_Log"] = result.tables.get("conversion_log", pd.DataFrame())
    sheets["Review_Items"] = result.tables.get("review_report", pd.DataFrame())
    sheets["Pipeline_Log"] = result.tables.get("pipeline_log", pd.DataFrame())

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        for name, frame in sheets.items():
            frame = _excel_safe(frame if len(frame.columns) else pd.DataFrame({"info": ["(none)"]}))
            frame.to_excel(writer, sheet_name=name, index=False)
            _style_sheet(writer.sheets[name], frame)
    return buf.getvalue()


# ----------------------------------------------------------------------
# 7. EDA
# ----------------------------------------------------------------------
def _subplots(nrows: int = 1, ncols: int = 1, figsize=(6, 4)):
    """Create a Figure outside pyplot's global registry so it is garbage-collected."""
    fig = Figure(figsize=figsize)
    return fig, fig.subplots(nrows, ncols)


def _style():
    sns.set_theme(style="whitegrid", context="notebook")
    plt.rcParams.update({"axes.titleweight": "bold", "axes.spines.top": False, "axes.spines.right": False})


def build_eda(df_original: pd.DataFrame, model_df: pd.DataFrame, numeric: list, categorical: list, cfg: PipelineConfig):
    """Return (figures, tables) describing the cleaned feature data."""
    _style()
    figures, tables = {}, {}

    missing = df_original.isna().mean().mul(100)
    missing = missing[missing > 0].sort_values().tail(25)
    if not missing.empty:
        fig, ax = _subplots(figsize=(8, max(2.6, 0.34 * len(missing) + 1.2)))
        ax.barh(missing.index.astype(str), missing.values, color=PRIMARY)
        ax.set_xlabel("Missing values (%)")
        ax.set_title("Missing values by column")
        fig.tight_layout()
        figures["Missing values"] = fig

    usable = [c for c in numeric if model_df[c].dropna().nunique() > 1]
    chosen = usable[: cfg.max_eda_columns]
    if chosen:
        ncols = 3
        nrows = int(np.ceil(len(chosen) / ncols))
        fig, axes = _subplots(nrows, ncols, figsize=(12, 3.0 * nrows))
        axes = np.atleast_1d(axes).ravel()
        for ax, col in zip(axes, chosen):
            values = model_df[col].dropna()
            ax.hist(values, bins=30, color=PRIMARY, alpha=0.85)
            ax.set_title(str(col)[:30])
            ax.set_ylabel("")
        for ax in axes[len(chosen):]:
            ax.axis("off")
        fig.suptitle("Numeric feature distributions", fontweight="bold", y=1.0)
        fig.tight_layout()
        figures["Numeric distributions"] = fig

    if len(usable) >= 2:
        subset = usable[:25]
        corr = model_df[subset].corr()
        size = min(11, max(5, 0.5 * len(subset) + 2))
        fig, ax = _subplots(figsize=(size, size * 0.85))
        sns.heatmap(
            corr, cmap="RdBu_r", center=0, vmin=-1, vmax=1, annot=len(subset) <= 10,
            fmt=".2f", linewidths=0.3, ax=ax, cbar_kws={"shrink": 0.8},
        )
        ax.set_title("Correlation between numeric features")
        fig.tight_layout()
        figures["Correlation heatmap"] = fig

    cats = [c for c in categorical if model_df[c].dropna().nunique() > 1][:6]
    if cats:
        ncols = 3
        nrows = int(np.ceil(len(cats) / ncols))
        fig, axes = _subplots(nrows, ncols, figsize=(13, 3.4 * nrows))
        axes = np.atleast_1d(axes).ravel()
        for ax, col in zip(axes, cats):
            counts = model_df[col].value_counts().head(8).sort_values()
            ax.barh(counts.index.astype(str).str.slice(0, 22), counts.values, color=PRIMARY, alpha=0.85)
            ax.set_title(str(col)[:30])
        for ax in axes[len(cats):]:
            ax.axis("off")
        fig.suptitle("Top categories", fontweight="bold", y=1.0)
        fig.tight_layout()
        figures["Categorical breakdown"] = fig

    if numeric:
        described = model_df[numeric].describe().T.round(3)
        described.insert(0, "column", described.index)
        tables["numeric_summary"] = described.reset_index(drop=True)

        outliers = []
        for col in numeric:
            values = model_df[col].dropna()
            if values.nunique() < 4:
                continue
            q1, q3 = values.quantile([0.25, 0.75])
            iqr = q3 - q1
            if iqr == 0:
                continue
            n_out = int(((values < q1 - 1.5 * iqr) | (values > q3 + 1.5 * iqr)).sum())
            outliers.append({"column": col, "iqr_outliers": n_out, "percent": round(100 * n_out / len(values), 2)})
        if outliers:
            tables["iqr_outliers"] = pd.DataFrame(outliers).sort_values("iqr_outliers", ascending=False, ignore_index=True)

    if categorical:
        tables["categorical_summary"] = pd.DataFrame(
            [
                {
                    "column": c,
                    "unique_values": int(model_df[c].nunique()),
                    "most_common": str(model_df[c].mode().iloc[0]) if model_df[c].notna().any() else "",
                    "most_common_share_pct": round(float(model_df[c].value_counts(normalize=True).iloc[0] * 100), 2)
                    if model_df[c].notna().any() else 0.0,
                }
                for c in categorical
            ]
        )

    return figures, tables


# ----------------------------------------------------------------------
# 8. Autoencoder back-ends
# ----------------------------------------------------------------------
def layer_sizes(input_dim: int):
    """128/64/32 for wide inputs; scaled down so narrow inputs keep a real bottleneck."""
    latent = max(2, min(32, input_dim // 2))
    h2 = min(64, max(2 * latent, 4))
    h1 = min(128, max(2 * h2, 8))
    return h1, h2, latent


class KerasAutoencoder:
    backend = "TensorFlow / Keras"

    def __init__(self, input_dim: int, seed: int):
        import tensorflow as tf
        from tensorflow.keras import Input, Model
        from tensorflow.keras.layers import Dense

        tf.keras.utils.set_random_seed(seed)
        h1, h2, latent = layer_sizes(input_dim)
        inputs = Input(shape=(input_dim,), name="input_layer")
        x = Dense(h1, activation="relu", name="encoder_layer_1")(inputs)
        x = Dense(h2, activation="relu", name="encoder_layer_2")(x)
        z = Dense(latent, activation="relu", name="latent_layer")(x)
        x = Dense(h2, activation="relu", name="decoder_layer_1")(z)
        x = Dense(h1, activation="relu", name="decoder_layer_2")(x)
        outputs = Dense(input_dim, activation="linear", name="output_layer")(x)
        self.model = Model(inputs, outputs, name="DataLens_AI_Autoencoder")
        self.model.compile(optimizer="adam", loss="mse")

    def fit(self, X, cfg: PipelineConfig, on_epoch=None) -> dict:
        from tensorflow.keras.callbacks import Callback, EarlyStopping

        class _Progress(Callback):
            def on_epoch_end(self, epoch, logs=None):
                if on_epoch:
                    on_epoch(epoch + 1, cfg.epochs, logs or {})

        history = self.model.fit(
            X, X,
            validation_split=cfg.validation_split,
            epochs=cfg.epochs,
            batch_size=cfg.batch_size,
            shuffle=True,
            callbacks=[
                EarlyStopping(monitor="val_loss", patience=cfg.patience, restore_best_weights=True),
                _Progress(),
            ],
            verbose=0,
        )
        return {
            "loss": [float(v) for v in history.history["loss"]],
            "val_loss": [float(v) for v in history.history.get("val_loss", [])],
        }

    def reconstruct(self, X):
        return self.model.predict(X, batch_size=1024, verbose=0)

    def save(self, directory: Path) -> Path:
        path = Path(directory) / "anomaly_autoencoder.keras"
        self.model.save(path)
        return path


class SklearnAutoencoder:
    """Fallback used only when TensorFlow is not installed."""

    backend = "scikit-learn MLP (TensorFlow not available)"

    def __init__(self, input_dim: int, seed: int):
        self.input_dim, self.seed = input_dim, seed
        self.model = None

    def fit(self, X, cfg: PipelineConfig, on_epoch=None) -> dict:
        from sklearn.neural_network import MLPRegressor

        h1, h2, latent = layer_sizes(self.input_dim)
        self.model = MLPRegressor(
            hidden_layer_sizes=(h1, h2, latent, h2, h1),
            activation="relu",
            solver="adam",
            batch_size=min(cfg.batch_size, len(X)),
            max_iter=cfg.epochs,
            early_stopping=True,
            validation_fraction=cfg.validation_split,
            n_iter_no_change=cfg.patience,
            random_state=self.seed,
        )
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            self.model.fit(X, X)
        loss = [float(v) for v in self.model.loss_curve_]
        if on_epoch:
            on_epoch(len(loss), cfg.epochs, {"loss": loss[-1]})
        return {"loss": loss, "val_loss": []}

    def reconstruct(self, X):
        return self.model.predict(X).astype(np.float32)

    def save(self, directory: Path) -> Path:
        path = Path(directory) / "anomaly_autoencoder.joblib"
        joblib.dump(self.model, path)
        return path


def make_autoencoder(input_dim: int, seed: int):
    """Prefer Keras; fall back to scikit-learn with an explicit warning."""
    try:
        return KerasAutoencoder(input_dim, seed), None
    except Exception as exc:  # noqa: BLE001
        note = (
            f"TensorFlow could not be used ({type(exc).__name__}). "
            "A scikit-learn neural-network autoencoder was used instead. "
            "Install TensorFlow to use the original Keras model."
        )
        return SklearnAutoencoder(input_dim, seed), note


# ----------------------------------------------------------------------
# 9. Scoring and explanations
# ----------------------------------------------------------------------
def _feature_groups(preprocessor, numeric: list, categorical: list):
    """Map each transformed model column back to its source column."""
    names = list(preprocessor.get_feature_names_out())
    owners = []
    for name in names:
        if name.startswith("numeric__"):
            owners.append(name[len("numeric__"):])
        else:
            raw = name[len("categorical__"):]
            matches = [c for c in categorical if raw.startswith(f"{c}_")]
            owners.append(max(matches, key=len) if matches else raw)
    order = list(dict.fromkeys(owners))
    groups = [[i for i, o in enumerate(owners) if o == col] for col in order]
    return order, groups


# ----------------------------------------------------------------------
# 10. The pipeline
# ----------------------------------------------------------------------
def _report(progress: Optional[ProgressFn], fraction: float, message: str):
    if progress:
        progress(min(max(fraction, 0.0), 1.0), message)


def run_pipeline(
    file_bytes: bytes,
    filename: str,
    cfg: Optional[PipelineConfig] = None,
    sheet_name=None,
    progress: Optional[ProgressFn] = None,
) -> PipelineResult:
    cfg = cfg or PipelineConfig()
    started = time.time()
    notes, log = [], []

    def step(name, detail):
        log.append({"step": name, "detail": detail})

    # --- Load -----------------------------------------------------------
    _report(progress, 0.03, "Loading and validating the dataset…")
    df = load_dataset(file_bytes, filename, sheet_name)
    if len(df) < cfg.min_rows:
        raise ValueError(
            f"The dataset has only {len(df)} rows. At least {cfg.min_rows} rows are needed "
            "to train and evaluate an anomaly model."
        )
    if len(df) < 200:
        notes.append("Small dataset (< 200 rows): anomaly scores will be less reliable.")
    step("Load", f"{len(df):,} rows × {df.shape[1]} columns from '{filename}'")

    # --- Profile --------------------------------------------------------
    _report(progress, 0.10, "Profiling data quality…")
    profile = profile_dataset(df)
    duplicate_rows = int(df.duplicated().sum())
    step("Profile", f"{duplicate_rows:,} duplicate rows; "
                    f"{int((profile['missing_count'] > 0).sum())} columns contain missing values")

    near_duplicates = compare_columns(df)

    # --- Clean ----------------------------------------------------------
    _report(progress, 0.18, "Detecting dates and converting numeric text…")
    temporal = detect_temporal_columns(df, dayfirst=cfg.dayfirst)
    converted, conversion_log, review_report = convert_numeric_text_columns(
        df, temporal, cfg.numeric_text_threshold
    )
    step("Convert numeric text", f"{len(conversion_log)} text columns converted to numbers; "
                                 f"{len(review_report)} values left for review")
    step("Detect dates", f"{len(temporal)} date/year columns detected")

    # --- Column plan ----------------------------------------------------
    _report(progress, 0.30, "Choosing model features…")
    plan = plan_columns(converted, temporal, cfg)
    model_df, numeric, categorical, date_info = build_model_frame(converted, plan, temporal, cfg)
    if model_df.shape[1] == 0:
        raise ValueError(
            "No usable feature columns were found. Every column was excluded as empty, "
            "constant, duplicate, or identifier-like. Check the 'Column plan' for details."
        )
    excluded = plan.loc[plan["role"] == "excluded", "column"].tolist()
    step("Feature selection", f"{len(numeric)} numeric + {len(categorical)} categorical features; "
                              f"{len(excluded)} columns excluded")

    # --- EDA ------------------------------------------------------------
    _report(progress, 0.40, "Generating exploratory analysis…")
    figures, eda_tables = build_eda(df, model_df, numeric, categorical, cfg)

    # --- Preprocess -----------------------------------------------------
    _report(progress, 0.50, "Splitting data and fitting the preprocessor…")
    positions = np.arange(len(model_df))
    train_pos, test_pos = train_test_split(positions, test_size=cfg.test_size, random_state=cfg.random_state)

    transformers = []
    if numeric:
        transformers.append((
            "numeric",
            Pipeline([("imputer", SimpleImputer(strategy="median")), ("scaler", StandardScaler())]),
            numeric,
        ))
    if categorical:
        transformers.append((
            "categorical",
            Pipeline([
                ("imputer", SimpleImputer(strategy="most_frequent")),
                ("encoder", OneHotEncoder(
                    handle_unknown="ignore",
                    min_frequency=cfg.min_category_frequency,
                    max_categories=cfg.max_onehot_categories,
                    sparse_output=False,
                )),
            ]),
            categorical,
        ))
    preprocessor = ColumnTransformer(transformers, remainder="drop")
    preprocessor.fit(model_df.iloc[train_pos])  # fit on training rows only
    X_all = np.asarray(preprocessor.transform(model_df), dtype=np.float32)
    if not np.isfinite(X_all).all():
        raise ValueError("Preprocessing produced non-finite values. Check the dataset for extreme numbers.")
    X_train = X_all[train_pos]
    step("Preprocess", f"{len(train_pos):,} training rows / {len(test_pos):,} held-out rows; "
                       f"{X_all.shape[1]} model inputs")

    # --- Train ----------------------------------------------------------
    _report(progress, 0.55, "Training the autoencoder…")
    autoencoder, backend_note = make_autoencoder(X_all.shape[1], cfg.random_state)
    if backend_note:
        notes.append(backend_note)
    np.random.seed(cfg.random_state)

    def on_epoch(epoch, total, logs):
        loss = logs.get("loss")
        text = f"Training the autoencoder… epoch {epoch}/{total}"
        if loss is not None:
            text += f" (loss {loss:.4f})"
        _report(progress, 0.55 + 0.33 * epoch / max(total, 1), text)

    history = autoencoder.fit(X_train, cfg, on_epoch=on_epoch)
    step("Train", f"{autoencoder.backend}; {len(history['loss'])} epochs completed")

    # --- Score ----------------------------------------------------------
    _report(progress, 0.90, "Scoring every row…")
    reconstructed = autoencoder.reconstruct(X_all)
    squared = np.square(X_all - reconstructed)
    errors = squared.mean(axis=1)
    threshold = float(np.percentile(errors[train_pos], cfg.threshold_percentile))
    flags = errors > threshold

    order, groups = _feature_groups(preprocessor, numeric, categorical)
    explanation = np.full(len(errors), "", dtype=object)
    flagged_pos = np.where(flags)[0]
    share_totals = np.zeros(len(order))
    if len(flagged_pos):
        per_col = np.column_stack([squared[flagged_pos][:, g].sum(axis=1) for g in groups])
        shares = per_col / np.maximum(per_col.sum(axis=1, keepdims=True), 1e-12)
        share_totals = shares.mean(axis=0)
        top = np.argsort(-shares, axis=1)[:, :3]
        for row_i, pos in enumerate(flagged_pos):
            explanation[pos] = ", ".join(f"{order[j]} ({shares[row_i, j]:.0%})" for j in top[row_i])

    split_label = np.full(len(errors), "train", dtype=object)
    split_label[test_pos] = "test"
    scored_pos = positions if cfg.score_all_rows else np.sort(test_pos)

    scored = df.iloc[scored_pos].copy(deep=True)
    new_columns = {
        "row_number": scored_pos + 1,
        "data_split": split_label[scored_pos],
        "reconstruction_error": errors[scored_pos],
        "anomaly_severity": errors[scored_pos] / max(threshold, 1e-12),
        "potential_anomaly": flags[scored_pos],
        "anomaly_status": np.where(flags[scored_pos], "Potential Anomaly", "Normal"),
        "top_contributing_features": explanation[scored_pos],
    }
    for name in new_columns:
        if name in scored.columns:
            scored = scored.rename(columns={name: f"{name}_input"})
    for name, values in new_columns.items():
        scored[name] = values
    leading = ["row_number"]
    scored = scored[leading + [c for c in scored.columns if c not in leading]]
    scored = scored.sort_values("reconstruction_error", ascending=False).reset_index(drop=True)

    anomalies = scored.loc[scored["potential_anomaly"]].copy().reset_index(drop=True)
    anomalies.insert(0, "anomaly_rank", np.arange(1, len(anomalies) + 1))
    step("Score", f"threshold = {cfg.threshold_percentile:g}th percentile of training error "
                  f"({threshold:.6f}); {len(anomalies):,} potential anomalies")

    # --- Training figures -----------------------------------------------
    _style()
    fig, ax = _subplots(figsize=(7, 3.6))
    ax.plot(range(1, len(history["loss"]) + 1), history["loss"], color=PRIMARY, label="Training loss")
    if history["val_loss"]:
        ax.plot(range(1, len(history["val_loss"]) + 1), history["val_loss"], color=DANGER, label="Validation loss")
    ax.set_xlabel("Epoch")
    ax.set_ylabel("Reconstruction loss (MSE)")
    ax.set_title("Training progress")
    ax.legend()
    fig.tight_layout()
    figures["Training loss"] = fig

    fig, ax = _subplots(figsize=(8, 3.8))
    log_err = np.log10(errors + 1e-12)
    ax.hist(log_err[~flags], bins=50, color=MUTED, alpha=0.9, label="Normal")
    ax.hist(log_err[flags], bins=50, color=DANGER, alpha=0.9, label="Potential anomaly")
    ax.axvline(np.log10(threshold + 1e-12), color="black", linestyle="--", label="Threshold")
    ax.set_xlabel("log10(reconstruction error)")
    ax.set_ylabel("Rows")
    ax.set_title("Reconstruction error distribution")
    ax.legend()
    fig.tight_layout()
    figures["Error distribution"] = fig

    contribution = (
        pd.DataFrame({"column": order, "share": share_totals})
        .sort_values("share", ascending=False, ignore_index=True)
        .head(10)
    )
    if len(flagged_pos):
        fig, ax = _subplots(figsize=(8, max(2.6, 0.38 * len(contribution) + 1.2)))
        view = contribution.iloc[::-1]
        ax.barh(view["column"].astype(str), view["share"] * 100, color=DANGER, alpha=0.85)
        ax.set_xlabel("Average share of anomaly error (%)")
        ax.set_title("Which columns drive the anomalies?")
        fig.tight_layout()
        figures["Anomaly drivers"] = fig

    # --- Summary --------------------------------------------------------
    test_flags = flags[test_pos]
    summary = {
        "Source file": filename,
        "Rows in dataset": len(df),
        "Columns in dataset": df.shape[1],
        "Features used by model": f"{len(numeric)} numeric + {len(categorical)} categorical",
        "Model inputs after encoding": X_all.shape[1],
        "Columns excluded": len(excluded),
        "Training rows": len(train_pos),
        "Held-out test rows": len(test_pos),
        "Rows scored": len(scored),
        "Model backend": autoencoder.backend,
        "Epochs completed": len(history["loss"]),
        "Final training loss": round(history["loss"][-1], 6),
        "Final validation loss": round(history["val_loss"][-1], 6) if history["val_loss"] else "n/a",
        "Threshold percentile": cfg.threshold_percentile,
        "Anomaly threshold (reconstruction error)": round(threshold, 6),
        "Potential anomalies": len(anomalies),
        "Potential anomaly rate (%)": round(100 * len(anomalies) / max(len(scored), 1), 2),
        "Held-out test anomalies": int(test_flags.sum()),
        "Held-out test anomaly rate (%)": round(100 * float(test_flags.mean()), 2) if len(test_flags) else 0.0,
        "Runtime (seconds)": round(time.time() - started, 1),
    }

    tables = {
        "profile": profile,
        "near_duplicate_columns": near_duplicates,
        "temporal_report": temporal,
        "conversion_log": conversion_log,
        "review_report": review_report,
        "column_plan": plan,
        "anomaly_drivers": contribution,
        "pipeline_log": pd.DataFrame(log),
        **eda_tables,
    }

    result = PipelineResult(
        filename=filename,
        config=cfg,
        original=df,
        scored=scored,
        anomalies=anomalies,
        summary=summary,
        tables=tables,
        figures=figures,
        history=history,
        warnings=notes,
    )

    # --- Export ---------------------------------------------------------
    _report(progress, 0.95, "Building the downloadable report…")
    result.excel_report = build_excel_report(result)

    metadata = {
        "pipeline_version": __version__,
        "model_backend": autoencoder.backend,
        "input_dimension": int(X_all.shape[1]),
        "anomaly_threshold": threshold,
        "threshold_percentile": cfg.threshold_percentile,
        "numeric_features": numeric,
        "categorical_features": categorical,
        "date_columns": date_info,
        "excluded_columns": excluded,
        "config": cfg.__dict__,
    }
    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp)
        model_path = autoencoder.save(tmp_path)
        joblib.dump(preprocessor, tmp_path / "anomaly_preprocessor.joblib")
        (tmp_path / "anomaly_metadata.json").write_text(json.dumps(metadata, indent=2, default=str), encoding="utf-8")
        bundle = io.BytesIO()
        with zipfile.ZipFile(bundle, "w", zipfile.ZIP_DEFLATED) as archive:
            for path in (model_path, tmp_path / "anomaly_preprocessor.joblib", tmp_path / "anomaly_metadata.json"):
                archive.write(path, arcname=path.name)
        result.artifacts_zip = bundle.getvalue()

    _report(progress, 1.0, "Done.")
    return result


__all__ = [
    "PipelineConfig",
    "PipelineResult",
    "list_excel_sheets",
    "load_dataset",
    "run_pipeline",
]
